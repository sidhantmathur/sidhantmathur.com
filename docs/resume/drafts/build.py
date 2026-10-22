"""Build the tailored resume drafts from one source of text.

Writes resume-<variant>.md and resume-<variant>.docx next to this file, for
each variant below. The layout copies public/resume.pdf (the current export of
the canonical resume): Letter, 0.5in side margins, Calibri 10pt, a 17pt name,
10.5pt section heads over a thin grey rule, company and location on one line,
title and dates on the next, bullets under them.

    python3 docs/resume/drafts/build.py
    # PDF (optional), from this folder:
    soffice --headless --convert-to pdf resume-revops.docx resume-delivery-ops.docx

Every fact here comes from docs/resume/resume-source.md. The one sentence that
comes from content/knowledge/ instead (the co-founder ran the 251-player event
on the ground) is flagged in README.md.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent

HEADER = {
    "name": "Sidhant Mathur",
    "lines": [
        "Toronto, ON | Canadian citizen | 613-795-9684 | hello@sidhantmathur.com",
        "linkedin.com/in/sidhantmathur | github.com/sidhantmathur | sidhantmathur.com",
    ],
}

EDUCATION = [
    ("B.S. in Business Administration — The University of Texas at Dallas", "2016 – 2019"),
    ("Software Engineering Immersive — General Assembly, Toronto", "2020 – 2021"),
]

# Citizenship rides in the header, where a screener checks work authorization
# first; a line of its own under the skills read as stranded.
OTHER = None

FREELANCE = {
    "org": "Freelance web development & consulting",
    "where": "Remote",
    "roles": [
        {
            "title": "Self-employed",
            "dates": "May 2019 – Oct 2022",
            "bullets": [
                "Advised a startup client on web-development strategy, saving $100k+ in agency fees.",
                "Built custom client solutions: an integrated web ordering experience and an affiliate-marketing site.",
            ],
        }
    ],
}

DELL = {
    "org": "Dell Technologies",
    "where": "Austin, TX",
    "roles": [
        {
            "title": "Marketing intern",
            "dates": "May 2018 – Aug 2018",
            "bullets": [
                "Built a propensity model (Azure ML) on Dell's historical customer database to find "
                "cross-sell targets for a high-margin software product. It surfaced ~20,000 qualified "
                "accounts, an estimated $129M pipeline (segment-level average deal size).",
            ],
        }
    ],
}

WFM_BULLETS = [
    "Owned headcount reporting and forecasting for a 3,000–4,000-person organization.",
    "Identified $1.2M in real-estate cost savings by analyzing lab and facility utilization data.",
]

VARIANTS = {
    # ------------------------------------------------------------------ RevOps
    "revops": {
        "label": "RevOps / GTM systems / sales ops",
        "summary": (
            "Sales operations specialist at Nokia. Builds the sales reporting, data models, and "
            "internal tools for global sales and finance teams, and owns each one through rollout "
            "and support. Co-founder & CTO of A Darle 20, a live two-sided marketplace "
            "(1,400+ users, 2,100+ bookings) built solo with an AI-agent-heavy workflow in Claude Code. "
            "Looking for a RevOps or GTM systems role at an AI-forward company."
        ),
        "experience": [
            {
                "org": "Nokia",
                "where": "Toronto, ON",
                "roles": [
                    {
                        "title": "Sales operations specialist (internal tools & automation)",
                        "dates": "Jun 2024 – Present",
                        "bullets": [
                            "Migrated all sales reporting from Salesforce Analytics to Power BI for 150+ "
                            "users. Owned the data model, the transformation logic (SQL/DAX), and the "
                            "global stakeholder rollout end to end.",
                            "Built a self-serve quarterly executive reporting app (Power Apps) used by 80+ "
                            "stakeholders across 7 regions. It eliminated a manual collection process that "
                            "consumed one finance analyst's entire workday, every day of quarter-end close.",
                            "Own every tool through its full lifecycle: turn vague requests from sales and "
                            "finance leadership into scoped requirements, then run executive demos, global "
                            "rollout, user support, and iteration across time zones.",
                        ],
                    },
                    {
                        "title": "Workforce management specialist",
                        "dates": "Oct 2022 – Jun 2024",
                        "bullets": WFM_BULLETS,
                    },
                ],
            },
            {
                "org": "A Darle 20 (“Let’s Go 20”) — adarle20.com",
                "where": "Toronto, ON",
                "roles": [
                    {
                        "title": "Co-founder & CTO (solo developer)",
                        "dates": "Aug 2025 – Present · launched Mar 2026",
                        "bullets": [
                            "Grew to 1,400+ registered users, 127 hosts, and 2,100+ bookings in the first "
                            "four months after launch, with 200–400 unique visitors daily.",
                            "Instrumented the funnel end to end (bookings, conversion, cancellations, "
                            "refunds, host activation) and use the data to steer product and monetization "
                            "decisions.",
                            "Integrated Stripe Connect for host payouts, platform fees, and OXXO cash "
                            "payments, designing payment flows for a Mexican customer base with low card "
                            "penetration, with automated refunds and transactional email.",
                            "Architected and shipped the entire marketplace as solo developer (TypeScript, "
                            "Next.js, Supabase/PostgreSQL, Stripe Connect, Vercel) with an AI-agent-heavy "
                            "workflow in Claude Code, keeping hands-on ownership of architecture, code "
                            "review, and release.",
                            "Run an autonomous multi-agent code-audit workflow (Claude Code) that reviews, "
                            "refactors, and regression-tests the production codebase.",
                            "Built the technology behind a flagship 251-player event at a 19,000-attendee "
                            "convention, rated 4.96/5 by attendees.",
                        ],
                    }
                ],
            },
            FREELANCE,
            DELL,
        ],
        "skills": [
            ("Sales reporting & data", "SQL, data modeling, DAX, Power BI, Salesforce, Excel/VBA, forecasting"),
            ("Internal tools & automation", "Power Apps / Power Automate, SharePoint, Python automation, ETL & data workflows"),
            ("AI-assisted development", "Claude Code, Anthropic API, LLM integration & agentic workflows, prompt engineering"),
            ("Engineering", "TypeScript, React, Next.js, Node.js, PostgreSQL, Supabase, Stripe, REST APIs"),
        ],
    },
    # ---------------------------------------------------------- Delivery ops
    "delivery-ops": {
        "label": "Delivery / professional services / deployment operations",
        "summary": (
            "3+ years in operations at Nokia: owned headcount reporting and forecasting for a "
            "3,000–4,000-person organization, then moved to sales operations, taking internal tools "
            "from vague request to global rollout and support across time zones. Co-founder & CTO of A Darle 20, a live marketplace: built its payment "
            "and refund flows, instrumented its funnel, and built the technology behind a 251-player "
            "live event. "
            "Looking for a delivery operations role at an AI-forward company."
        ),
        "experience": [
            {
                "org": "Nokia",
                "where": "Toronto, ON",
                "roles": [
                    {
                        "title": "Sales operations specialist (internal tools & automation)",
                        "dates": "Jun 2024 – Present",
                        "bullets": [
                            "Own every internal tool through its full lifecycle: turn vague requests from "
                            "sales and finance leadership into scoped requirements, then run executive "
                            "demos, global rollout, user support, and iteration across time zones.",
                            "Built the self-serve app (Power Apps) that runs the quarterly executive "
                            "reporting cadence for 80+ stakeholders across 7 regions. It eliminated a manual "
                            "collection process that took one finance analyst's entire workday, every day "
                            "of quarter-end close.",
                            "Migrated all sales reporting from Salesforce Analytics to Power BI for 150+ "
                            "users, owning the data model, transformation logic (SQL/DAX), and global "
                            "rollout.",
                        ],
                    },
                    {
                        "title": "Workforce management specialist",
                        "dates": "Oct 2022 – Jun 2024",
                        "bullets": WFM_BULLETS,
                    },
                ],
            },
            {
                "org": "A Darle 20 (“Let’s Go 20”) — adarle20.com",
                "where": "Toronto, ON",
                "roles": [
                    {
                        "title": "Co-founder & CTO (solo developer)",
                        "dates": "Aug 2025 – Present · launched Mar 2026",
                        "bullets": [
                            "Built the technology behind a flagship 251-player event at a 19,000-attendee "
                            "convention, rated 4.96/5 by attendees; a co-founder ran the event on the ground.",
                            "Integrated Stripe Connect for host payouts, platform fees, and OXXO cash "
                            "payments, designing payment flows for a Mexican customer base with low card "
                            "penetration, with automated refunds and transactional email.",
                            "Instrumented the funnel end to end (bookings, conversion, cancellations, "
                            "refunds, host activation) and use the data to steer product and monetization "
                            "decisions.",
                            "Grew to 1,400+ registered users, 127 hosts, and 2,100+ bookings in the first "
                            "four months after launch, with 200–400 unique visitors daily.",
                            "Architected and shipped the entire marketplace as solo developer (TypeScript, "
                            "Next.js, Supabase/PostgreSQL, Stripe Connect, Vercel) with an AI-agent-heavy "
                            "workflow in Claude Code; a multi-agent code-audit workflow reviews, refactors, "
                            "and regression-tests the production codebase.",
                        ],
                    }
                ],
            },
            FREELANCE,
            DELL,
        ],
        "skills": [
            ("Reporting & forecasting", "SQL, data modeling, Power BI, DAX, Salesforce, Excel/VBA, forecasting"),
            ("Internal tools & automation", "Power Apps / Power Automate, SharePoint, Python automation, ETL & data workflows"),
            ("Payments & product", "Stripe Connect, Supabase/PostgreSQL, TypeScript, Next.js, React, Node.js, REST APIs"),
            ("AI-assisted development", "Claude Code, Anthropic API, LLM integration & agentic workflows, prompt engineering"),
        ],
    },
}


# --------------------------------------------------------------------------- markdown

def to_markdown(v: dict) -> str:
    out = [f"# {HEADER['name']}", "", HEADER["lines"][0].replace(" | ", " · "),
           HEADER["lines"][1].replace(" | ", " · "), "", "## Summary", "", v["summary"], "",
           "## Experience", ""]
    for job in v["experience"]:
        out += [f"### {job['org']} — {job['where']}", ""]
        for role in job["roles"]:
            out += [f"**{role['title']} · {role['dates']}**", ""]
            out += [f"- {b}" for b in role["bullets"]]
            out.append("")
    out += ["## Education", ""]
    out += [f"- {what}, {when}" for what, when in EDUCATION]
    out += ["", "## Technical skills", ""]
    out += [f"- **{k}:** {val}" for k, val in v["skills"]]
    if OTHER:
        out += ["", "## Other", "", f"- {OTHER}", ""]
    else:
        out += [""]
    return "\n".join(out)


# --------------------------------------------------------------------------- docx

FONT = "Calibri"
BODY = Pt(10)
TEXT_WIDTH = Inches(7.5)  # 8.5in page minus 0.5in each side


def _font(run, size=BODY, bold=False, italic=False, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    run.font.size = size
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def _spacing(p, before=0, after=0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0


def _rule_below(p):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    bdr.append(bottom)
    ppr.append(bdr)


def _split_line(doc, left, right, left_style, before=0, after=0):
    p = doc.add_paragraph()
    _spacing(p, before, after)
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    _font(p.add_run(left), **left_style)
    _font(p.add_run("\t" + right))
    return p


def _bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    _spacing(p, 0, 1)
    pf = p.paragraph_format
    pf.left_indent = Inches(0.22)
    pf.first_line_indent = Inches(-0.14)
    _font(p.add_run(text))


def _heading(doc, text):
    p = doc.add_paragraph()
    _spacing(p, 7, 3)
    _rule_below(p)
    _font(p.add_run(text), size=Pt(10.5), bold=True, color=RGBColor(0x1A, 0x1A, 0x1A))


def to_docx(v: dict, path: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.5)
    sec.top_margin = sec.bottom_margin = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = BODY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, 0, 2)
    _font(p.add_run(HEADER["name"]), size=Pt(17), bold=True)
    for line in HEADER["lines"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(p, 0, 0)
        _font(p.add_run(line))

    _heading(doc, "Summary")
    p = doc.add_paragraph()
    _spacing(p, 0, 0)
    _font(p.add_run(v["summary"]))

    _heading(doc, "Experience")
    for i, job in enumerate(v["experience"]):
        _split_line(doc, job["org"], job["where"], {"bold": True}, before=0 if i == 0 else 5)
        for role in job["roles"]:
            _split_line(doc, role["title"], role["dates"], {"italic": True}, after=1)
            for b in role["bullets"]:
                _bullet(doc, b)

    _heading(doc, "Education")
    for what, when in EDUCATION:
        _split_line(doc, what, when, {})

    _heading(doc, "Technical skills")
    for k, val in v["skills"]:
        p = doc.add_paragraph()
        _spacing(p, 0, 1)
        _font(p.add_run(f"{k}: "), bold=True)
        _font(p.add_run(val))
    if OTHER:
        p = doc.add_paragraph()
        _spacing(p, 2, 0)
        _font(p.add_run(OTHER), bold=True)

    doc.core_properties.author = "Sidhant Mathur"
    doc.core_properties.title = f"Sidhant Mathur — resume ({v['label']})"
    doc.save(path)


if __name__ == "__main__":
    for key, variant in VARIANTS.items():
        (HERE / f"resume-{key}.md").write_text(to_markdown(variant), encoding="utf-8")
        to_docx(variant, HERE / f"resume-{key}.docx")
        print("wrote", f"resume-{key}.md", f"resume-{key}.docx")
